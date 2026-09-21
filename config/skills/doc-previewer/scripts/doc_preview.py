# -*- coding: utf-8 -*-
import os
import sys
import argparse
import subprocess
from pathlib import Path

def preview_pdf(file_path: str, output_dir: str, dpi: int = 150) -> dict:
    import pymupdf as fitz
    doc = fitz.open(file_path)
    os.makedirs(output_dir, exist_ok=True)
    images = []
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=mat, alpha=False)
        img_name = f"page_{i+1:02d}.png"
        out_p = os.path.join(output_dir, img_name)
        pix.save(out_p)
        images.append(out_p.replace("\\", "/"))
    doc.close()
    return {"total_pages": len(images), "images": images, "output_dir": output_dir.replace("\\", "/")}

def preview_docx(file_path: str, output_file: str) -> dict:
    import mammoth
    with open(file_path, "rb") as f:
        res = mammoth.convert_to_html(f)
    
    filename = Path(file_path).name
    styled_html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Văn bản: {filename}</title>
<link href="https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&family=Lora:ital,wght@0,400;0,600;1,400&display=swap" rel="stylesheet">
<style>
:root {{
    --bg-page: #f1f5f9;
    --paper: #ffffff;
    --text-main: #1e293b;
    --primary: #1e3a8a;
    --border: #cbd5e1;
}}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{
    font-family: 'Be Vietnam Pro', -apple-system, sans-serif;
    background: var(--bg-page);
    color: var(--text-main);
    line-height: 1.7;
    padding: 30px 16px;
    display: flex;
    justify-content: center;
}}
.paper-container {{
    max-width: 900px;
    width: 100%;
    background: var(--paper);
    border: 1px solid var(--border);
    border-radius: 12px;
    padding: 50px 60px;
    box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.1), 0 8px 10px -6px rgba(0, 0, 0, 0.1);
}}
.doc-header-banner {{
    background: #eff6ff;
    border-left: 5px solid #2563eb;
    padding: 16px 20px;
    border-radius: 6px;
    margin-bottom: 30px;
    display: flex;
    align-items: center;
    justify-content: space-between;
}}
.doc-title {{
    font-size: 18px;
    font-weight: 700;
    color: #1e3a8a;
}}
.doc-badge {{
    background: #2563eb;
    color: #ffffff;
    padding: 4px 10px;
    border-radius: 4px;
    font-size: 13px;
    font-weight: 600;
}}
.content-body {{
    font-size: 17px;
    color: #334155;
}}
.content-body h1, .content-body h2, .content-body h3 {{
    color: #0f172a;
    margin-top: 28px;
    margin-bottom: 12px;
    font-weight: 800;
    line-height: 1.3;
}}
.content-body h1 {{ font-size: 26px; border-bottom: 2px solid #e2e8f0; padding-bottom: 8px; color: #1e3a8a; }}
.content-body h2 {{ font-size: 22px; color: #1d4ed8; }}
.content-body h3 {{ font-size: 19px; color: #047857; }}
.content-body p {{ margin-bottom: 16px; text-align: justify; }}
.content-body table {{
    width: 100%;
    border-collapse: collapse;
    margin: 24px 0;
    font-size: 15px;
}}
.content-body th, .content-body td {{
    border: 1px solid #cbd5e1;
    padding: 10px 14px;
    text-align: left;
}}
.content-body th {{
    background: #f8fafc;
    font-weight: 700;
    color: #0f172a;
}}
.content-body tr:nth-child(even) {{ background: #f8fafc; }}
.content-body blockquote {{
    border-left: 4px solid #3b82f6;
    background: #f8fafc;
    padding: 12px 18px;
    margin: 18px 0;
    font-style: italic;
    color: #475569;
}}
.content-body ul, .content-body ol {{
    margin: 16px 0 16px 28px;
    display: flex;
    flex-direction: column;
    gap: 8px;
}}
</style>
</head>
<body>
<div class="paper-container">
    <div class="doc-header-banner">
        <div class="doc-title">📄 Tài Liệu Văn Bản: {filename}</div>
        <span class="doc-badge">DOCX PREVIEW</span>
    </div>
    <div class="content-body">
        {res.value}
    </div>
</div>
</body>
</html>"""
    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as f:
        f.write(styled_html)
    return {"output_file": output_file.replace("\\", "/")}

def preview_pptx(file_path: str, output_file: str) -> dict:
    from pptx import Presentation
    prs = Presentation(file_path)
    filename = Path(file_path).name
    
    slides_html = []
    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = p.text.strip()
                    if txt:
                        slide_texts.append(txt)
                        
        title = slide_texts[0] if slide_texts else f"Trang {slide_num}"
        bullets = slide_texts[1:] if len(slide_texts) > 1 else []
        
        items_markup = "".join([f"<li class='bullet-item'>{b}</li>" for b in bullets])
        slides_html.append(f"""
        <div class="slide {'active' if slide_num == 1 else ''}" id="slide-{slide_num}">
            <div class="slide-header">
                <span class="slide-tag">SLIDE {slide_num} / {len(prs.slides)}</span>
                <h2 class="slide-title">{title}</h2>
            </div>
            <div class="slide-content">
                <ul class="bullet-list">
                    {items_markup if items_markup else "<p class='empty-note'>[Nội dung trình bày đồ họa / bảng biểu]</p>"}
                </ul>
            </div>
            <div class="slide-footer">
                <span>{filename}</span>
                <span>Trang {slide_num}</span>
            </div>
        </div>
        """)
        
    all_slides = "\n".join(slides_html)
    styled_html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Slide Presentation: {filename}</title>
<link href="https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;600;700;800;900&display=swap" rel="stylesheet">
<style>
* {{ box-sizing: border-box; margin: 0; padding: 0; font-family: 'Be Vietnam Pro', sans-serif; }}
body {{
    background: #0f172a;
    color: #1e293b;
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: center;
    min-height: 100vh;
    overflow: hidden;
}}
.nav-bar {{
    position: fixed;
    top: 0; left: 0; right: 0;
    height: 54px;
    background: rgba(15, 23, 42, 0.95);
    border-bottom: 1px solid rgba(255,255,255,0.1);
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 0 24px;
    z-index: 100;
}}
.nav-title {{ color: #ffffff; font-size: 16px; font-weight: 700; }}
.nav-controls {{ display: flex; gap: 12px; align-items: center; }}
.btn {{
    background: #1e293b;
    color: #ffffff;
    border: 1px solid #334155;
    padding: 6px 14px;
    border-radius: 6px;
    font-weight: 700;
    cursor: pointer;
    font-size: 14px;
}}
.btn:hover {{ background: #2563eb; border-color: #3b82f6; }}
.btn:disabled {{ opacity: 0.4; cursor: not-allowed; }}
.counter {{ color: #f8fafc; font-weight: 800; padding: 4px 10px; background: #334155; border-radius: 12px; }}

.stage {{
    width: 92vw;
    max-width: 1280px;
    aspect-ratio: 16 / 9;
    margin-top: 54px;
    position: relative;
    background: #ffffff;
    border-radius: 12px;
    box-shadow: 0 25px 50px -12px rgba(0,0,0,0.6);
    overflow: hidden;
}}
.slide {{
    position: absolute;
    top: 0; left: 0; width: 100%; height: 100%;
    display: flex; flex-direction: column;
    opacity: 0; visibility: hidden;
    transition: opacity 0.3s ease;
    background: #ffffff;
}}
.slide.active {{ opacity: 1; visibility: visible; z-index: 10; }}
.slide-header {{
    background: #1e3a8a;
    color: #ffffff;
    padding: 20px 36px;
    border-bottom: 4px solid #22c55e;
}}
.slide-tag {{ background: #f59e0b; color: #0f172a; padding: 4px 10px; border-radius: 4px; font-weight: 800; font-size: 14px; margin-bottom: 8px; display: inline-block; }}
.slide-title {{ font-size: 26px; font-weight: 800; }}
.slide-content {{
    flex: 1;
    padding: 32px 40px;
    overflow-y: auto;
    background: #f8fafc;
}}
.bullet-list {{ list-style: none; display: flex; flex-direction: column; gap: 14px; }}
.bullet-item {{
    font-size: 19px;
    line-height: 1.6;
    color: #1e293b;
    padding: 12px 18px;
    background: #ffffff;
    border-radius: 8px;
    border-left: 4px solid #2563eb;
    box-shadow: 0 2px 4px rgba(0,0,0,0.05);
}}
.slide-footer {{
    padding: 10px 36px;
    background: #e2e8f0;
    display: flex; justify-content: space-between;
    font-size: 14px; font-weight: 600; color: #64748b;
}}
</style>
</head>
<body>
    <div class="nav-bar">
        <div class="nav-title">📊 {filename}</div>
        <div class="nav-controls">
            <button class="btn" id="pBtn" onclick="prev()">◀ Trước</button>
            <span class="counter" id="cnt">1 / {len(prs.slides)}</span>
            <button class="btn" id="nBtn" onclick="next()">Sau ▶</button>
            <button class="btn" onclick="document.documentElement.requestFullscreen()" style="background:#059669;">⛶ Toàn Màn Hình</button>
        </div>
    </div>
    <div class="stage">
        {all_slides}
    </div>
    <script>
        let cur = 1;
        const total = {len(prs.slides)};
        function update() {{
            for(let i=1; i<=total; i++) {{
                const s = document.getElementById('slide-'+i);
                if(s) s.classList.toggle('active', i === cur);
            }}
            document.getElementById('cnt').innerText = cur + ' / ' + total;
            document.getElementById('pBtn').disabled = (cur === 1);
            document.getElementById('nBtn').disabled = (cur === total);
        }}
        function next() {{ if(cur < total) {{ cur++; update(); }} }}
        function prev() {{ if(cur > 1) {{ cur--; update(); }} }}
        document.addEventListener('keydown', e => {{
            if(e.key === 'ArrowRight' || e.key === ' ') next();
            if(e.key === 'ArrowLeft') prev();
        }});
        update();
    </script>
</body>
</html>"""
    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as f:
        f.write(styled_html)
    return {"output_file": output_file.replace("\\", "/")}

def preview_xlsx(file_path: str, output_file: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(file_path, data_only=True)
    filename = Path(file_path).name
    
    sheets_markup = []
    tabs_markup = []
    
    for s_idx, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        is_active = "active" if s_idx == 0 else ""
        tabs_markup.append(f'<button class="tab-btn {is_active}" onclick="showSheet({s_idx})">{sheet_name}</button>')
        
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            table_html = "<p style='padding:20px;'>Bảng tính trống.</p>"
        else:
            header = rows[0]
            header_html = "".join([f"<th>{c or ''}</th>" for c in header])
            tbody_rows = []
            for r in rows[1:]:
                if any(r):
                    tds = "".join([f"<td>{c if c is not None else ''}</td>" for c in r])
                    tbody_rows.append(f"<tr>{tds}</tr>")
            table_html = f"""<table><thead><tr>{header_html}</tr></thead><tbody>{''.join(tbody_rows)}</tbody></table>"""
            
        sheets_markup.append(f"""
        <div class="sheet-container {'active' if s_idx==0 else ''}" id="sheet-{s_idx}">
            <div class="table-wrapper">{table_html}</div>
        </div>
        """)
        
    styled_html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Excel Preview: {filename}</title>
<link href="https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700&display=swap" rel="stylesheet">
<style>
* {{ box-sizing: border-box; margin: 0; padding: 0; font-family: 'Be Vietnam Pro', sans-serif; }}
body {{ background: #f1f5f9; color: #1e293b; padding: 24px; }}
.header {{ background: #15803d; color: #ffffff; padding: 16px 24px; border-radius: 8px 8px 0 0; display: flex; justify-content: space-between; align-items: center; }}
.tabs {{ display: flex; gap: 4px; background: #e2e8f0; padding: 8px 12px 0; border-bottom: 2px solid #cbd5e1; }}
.tab-btn {{ border: none; background: #cbd5e1; padding: 8px 16px; font-weight: 700; border-radius: 6px 6px 0 0; cursor: pointer; }}
.tab-btn.active {{ background: #ffffff; color: #15803d; border-top: 3px solid #15803d; }}
.sheet-container {{ display: none; background: #ffffff; padding: 20px; border-radius: 0 0 8px 8px; box-shadow: 0 4px 6px rgba(0,0,0,0.05); }}
.sheet-container.active {{ display: block; }}
.table-wrapper {{ overflow-x: auto; max-height: 75vh; }}
table {{ border-collapse: collapse; width: 100%; font-size: 14px; }}
th, td {{ border: 1px solid #cbd5e1; padding: 8px 12px; text-align: left; }}
th {{ background: #f8fafc; font-weight: 700; position: sticky; top: 0; }}
tr:nth-child(even) {{ background: #f8fafc; }}
</style>
</head>
<body>
    <div class="header">
        <h2>📊 Bảng Tính: {filename}</h2>
        <span>EXCEL PREVIEW</span>
    </div>
    <div class="tabs">
        {''.join(tabs_markup)}
    </div>
    {''.join(sheets_markup)}
    <script>
        function showSheet(idx) {{
            document.querySelectorAll('.tab-btn').forEach((b, i) => b.classList.toggle('active', i === idx));
            document.querySelectorAll('.sheet-container').forEach((c, i) => c.classList.toggle('active', i === idx));
        }}
    </script>
</body>
</html>"""
    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as f:
        f.write(styled_html)
    return {"output_file": output_file.replace("\\", "/")}

def preview_video(file_path: str, output_file: str) -> dict:
    p = Path(file_path).resolve()
    filename = p.name
    size_mb = f"{p.stat().st_size / (1024*1024):.2f} MB" if p.exists() else "Unknown"
    file_uri = p.as_uri()
    out_dir = Path(output_file).resolve().parent
    try:
        rel_src = os.path.relpath(p, out_dir).replace("\\", "/")
    except ValueError:
        rel_src = file_uri
    
    http_src = f"http://localhost:3050/{p.name}" if "public" in str(p).replace("\\", "/") else ""

    import base64
    b64_uri = ""
    compact_candidate = p.parent / f"{p.stem}_compact.mp4"
    target_p = compact_candidate if compact_candidate.exists() else p
    if target_p.exists() and target_p.stat().st_size < 15 * 1024 * 1024:
        try:
            b64_str = base64.b64encode(target_p.read_bytes()).decode('ascii')
            b64_uri = f"data:video/mp4;base64,{b64_str}"
        except Exception:
            b64_uri = ""

    styled_html = f"""<!DOCTYPE html>
<html lang="vi" class="light">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Video Studio: {filename}</title>
<script src="https://www.gstatic.com/antigravity/web/dev/tailwindcss.min.js"></script>
<link href="https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&family=JetBrains+Mono:wght@500;700&display=swap" rel="stylesheet">
<style>
  :root {{
    --bg-card: #ffffff;
    --border-color: #e2e8f0;
    --text-main: #0f172a;
    --text-muted: #64748b;
    --primary: #0284c7;
  }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{
    font-family: 'Be Vietnam Pro', -apple-system, sans-serif;
    background: transparent;
    color: var(--text-main);
    display: flex;
    justify-content: center;
    padding: 12px;
  }}
  .glass-card {{
    background: rgba(255, 255, 255, 0.96);
    backdrop-filter: blur(16px);
    border: 1.5px solid #e2e8f0;
    border-radius: 18px;
    box-shadow: 0 12px 35px -8px rgba(0, 0, 0, 0.12), 0 4px 12px -2px rgba(0, 0, 0, 0.05);
    width: 100%;
    max-width: 520px;
    overflow: hidden;
  }}
  .video-stage {{
    position: relative;
    background: #000000;
    width: 100%;
    aspect-ratio: 9 / 16;
    max-height: 520px;
    display: flex;
    align-items: center;
    justify-content: center;
    cursor: pointer;
  }}
  video {{
    width: 100%;
    height: 100%;
    object-fit: contain;
    outline: none;
  }}
  .big-play-btn {{
    position: absolute;
    width: 72px;
    height: 72px;
    border-radius: 50%;
    background: rgba(2, 132, 199, 0.92);
    border: 3px solid #ffffff;
    box-shadow: 0 0 30px rgba(2, 132, 199, 0.6);
    display: flex;
    align-items: center;
    justify-content: center;
    color: white;
    cursor: pointer;
    transition: transform 0.2s, background 0.2s;
  }}
  .big-play-btn:hover {{
    transform: scale(1.08);
    background: #0284c7;
  }}
  .mono-badge {{
    font-family: 'JetBrains Mono', monospace;
  }}
</style>
</head>
<body>
  <div class="glass-card">
    <!-- Header Banner -->
    <div style="padding: 14px 18px; border-bottom: 1px solid #f1f5f9; display: flex; align-items: center; justify-content: space-between; background: linear-gradient(180deg, #fafafa 0%, #ffffff 100%);">
      <div style="display: flex; align-items: center; gap: 8px;">
        <span style="width: 10px; height: 10px; background: #0284c7; border-radius: 50%; display: inline-block;"></span>
        <div>
          <h1 style="font-size: 14px; font-weight: 800; color: #0369a1; line-height: 1.2;">ANTIGRAVITY 2.0 VIDEO STUDIO</h1>
          <p class="mono-badge" style="font-size: 11px; color: #64748b; font-weight: 600;">{filename} • {size_mb}</p>
        </div>
      </div>
      <span style="font-size: 11px; font-weight: 800; background: #e0f2fe; color: #0284c7; padding: 3px 8px; border-radius: 9999px;">STUDIO PRO</span>
    </div>

    <!-- Video Stage -->
    <div class="video-stage" onclick="togglePlay()">
      <video id="vidPlayer" playsinline preload="auto">
        {f'<source src="{b64_uri}" type="video/mp4">' if b64_uri else ''}
        <source src="{rel_src}" type="video/mp4">
        {f'<source src="{http_src}" type="video/mp4">' if http_src else ''}
        <source src="{file_uri}" type="video/mp4">
        Trình duyệt không hỗ trợ xem video trực tiếp.
      </video>
      <div id="bigPlayBtn" class="big-play-btn">
        <svg width="32" height="32" viewBox="0 0 24 24" fill="currentColor" style="margin-left: 4px;">
          <polygon points="5 3 19 12 5 21 5 3"></polygon>
        </svg>
      </div>
    </div>

    <!-- Interactive Progress Scrub Bar -->
    <div style="padding: 8px 18px 0; background: #ffffff;">
      <input type="range" id="scrubber" min="0" max="100" value="0" step="0.1" style="width: 100%; accent-color: #0284c7; cursor: pointer; height: 6px;">
      <div style="display: flex; justify-content: space-between; align-items: center; margin-top: 4px; font-size: 12px; color: #64748b;" class="mono-badge">
        <span id="timeCurrent" style="font-weight: 700; color: #0f172a;">00:00</span>
        <span id="frameBadge" style="font-size: 10px; background: #f1f5f9; padding: 2px 6px; border-radius: 4px;">F: 0 / 25 FPS</span>
        <span id="timeDuration">00:00</span>
      </div>
    </div>

    <!-- Primary Control Bar -->
    <div style="padding: 12px 18px; display: flex; flex-direction: column; gap: 10px;">
      <!-- Action Buttons Row 1: Play, Skip, Frame-Step -->
      <div style="display: flex; align-items: center; justify-content: space-between; gap: 8px;">
        <button id="btnPlay" onclick="togglePlay()" style="flex: 1; padding: 9px 14px; background: #0284c7; color: white; font-weight: 800; font-size: 13px; border: none; border-radius: 8px; cursor: pointer; display: flex; align-items: center; justify-content: center; gap: 6px; box-shadow: 0 4px 12px rgba(2, 132, 199, 0.25);">
          <span id="playIcon">▶</span>
          <span id="playText">PHÁT VIDEO</span>
        </button>

        <!-- Precision Frame Step Buttons -->
        <div style="display: flex; gap: 4px;">
          <button onclick="stepFrame(-1)" title="Lùi 1 Frame (0.04s) [Phím [ ]" style="padding: 8px 10px; background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 6px; font-size: 11px; font-weight: 700; color: #334155; cursor: pointer;">
            ◀ 1F
          </button>
          <button onclick="stepFrame(1)" title="Tiến 1 Frame (0.04s) [Phím ] ]" style="padding: 8px 10px; background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 6px; font-size: 11px; font-weight: 700; color: #334155; cursor: pointer;">
            1F ▶
          </button>
        </div>

        <!-- Speed Selector -->
        <select id="speedSelect" onchange="setSpeed(this.value)" style="padding: 8px; background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 6px; font-size: 11px; font-weight: 700; color: #334155; cursor: pointer; outline: none;">
          <option value="0.5">0.5x</option>
          <option value="1.0" selected>1.0x</option>
          <option value="1.25">1.25x</option>
          <option value="1.5">1.5x</option>
          <option value="2.0">2.0x</option>
        </select>
      </div>

      <!-- Action Buttons Row 2: Snapshot, Mute, PiP, Fullscreen -->
      <div style="display: flex; align-items: center; justify-content: space-between; gap: 6px; font-size: 11px;">
        <button onclick="captureSnapshot()" title="Chụp ảnh khung hình hiện tại" style="padding: 7px 10px; background: #f0fdf4; border: 1px solid #bbf7d0; color: #15803d; border-radius: 6px; font-weight: 700; cursor: pointer; display: flex; align-items: center; gap: 4px;">
          📸 Chụp ảnh frame
        </button>

        <button onclick="toggleMute()" id="btnMute" style="padding: 7px 10px; background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 6px; font-weight: 600; color: #475569; cursor: pointer;">
          🔊 Tiếng
        </button>

        <button onclick="togglePiP()" title="Chế độ Cửa sổ Nổi (PiP)" style="padding: 7px 10px; background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 6px; font-weight: 600; color: #475569; cursor: pointer;">
          📺 PiP
        </button>

        <button onclick="toggleFullScreen()" title="Toàn màn hình [Phím F]" style="padding: 7px 10px; background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 6px; font-weight: 600; color: #475569; cursor: pointer;">
          ⛶ Toàn màn hình
        </button>
      </div>
    </div>

    <!-- Footer Specs Legend -->
    <div style="background: #f8fafc; border-top: 1px solid #f1f5f9; padding: 10px 18px; font-size: 11px; color: #64748b; line-height: 1.4;">
      <div style="display: flex; justify-content: space-between; align-items: center;">
        <span>🎯 <b>Phím tắt:</b> Space (Play/Pause) • ◄/► (Tua 3s) • [/] (1 Frame)</span>
        <a href="{rel_src}" download style="color: #0284c7; font-weight: 800; text-decoration: none;">Tải file MP4 ⬇️</a>
      </div>
    </div>
  </div>

  <script>
    const vid = document.getElementById('vidPlayer');
    const scrubber = document.getElementById('scrubber');
    const timeCurrent = document.getElementById('timeCurrent');
    const timeDuration = document.getElementById('timeDuration');
    const frameBadge = document.getElementById('frameBadge');
    const bigPlayBtn = document.getElementById('bigPlayBtn');
    const playText = document.getElementById('playText');
    const playIcon = document.getElementById('playIcon');
    const btnMute = document.getElementById('btnMute');
    const FPS = 25.0;

    function formatTime(sec) {{
      if (isNaN(sec)) return "00:00";
      const m = Math.floor(sec / 60);
      const s = Math.floor(sec % 60);
      return String(m).padStart(2, '0') + ":" + String(s).padStart(2, '0');
    }}

    function togglePlay() {{
      if (vid.paused) {{
        vid.play().then(() => {{
          bigPlayBtn.style.display = 'none';
          playText.textContent = 'TẠM DỪNG';
          playIcon.textContent = '❚❚';
        }}).catch(() => {{
          vid.muted = true;
          vid.play().then(() => {{
            bigPlayBtn.style.display = 'none';
            playText.textContent = 'TẠM DỪNG (MUTE)';
            playIcon.textContent = '❚❚';
          }});
        }});
      }} else {{
        vid.pause();
        bigPlayBtn.style.display = 'flex';
        playText.textContent = 'PHÁT VIDEO';
        playIcon.textContent = '▶';
      }}
    }}

    vid.addEventListener('timeupdate', () => {{
      if (vid.duration) {{
        scrubber.value = (vid.currentTime / vid.duration) * 100;
        timeCurrent.textContent = formatTime(vid.currentTime);
        const f = Math.round(vid.currentTime * FPS);
        const totalF = Math.round(vid.duration * FPS);
        frameBadge.textContent = `F: ${{f}}/${{totalF}} @ ${{FPS}}FPS`;
      }}
    }});

    vid.addEventListener('loadedmetadata', () => {{
      timeDuration.textContent = formatTime(vid.duration);
      const totalF = Math.round(vid.duration * FPS);
      frameBadge.textContent = `F: 0/${{totalF}} @ ${{FPS}}FPS`;
    }});

    vid.addEventListener('ended', () => {{
      bigPlayBtn.style.display = 'flex';
      playText.textContent = 'PHÁT LẠI';
      playIcon.textContent = '↻';
    }});

    scrubber.addEventListener('input', () => {{
      if (vid.duration) {{
        vid.currentTime = (scrubber.value / 100) * vid.duration;
      }}
    }});

    function stepFrame(delta) {{
      vid.pause();
      vid.currentTime = Math.max(0, Math.min(vid.duration, vid.currentTime + delta * (1.0 / FPS)));
      bigPlayBtn.style.display = 'flex';
      playText.textContent = 'PHÁT VIDEO';
      playIcon.textContent = '▶';
    }}

    function setSpeed(s) {{
      vid.playbackRate = parseFloat(s);
    }}

    function toggleMute() {{
      vid.muted = !vid.muted;
      btnMute.textContent = vid.muted ? '🔇 Tắt tiếng' : '🔊 Tiếng';
    }}

    function togglePiP() {{
      if (document.pictureInPictureElement) {{
        document.exitPictureInPicture().catch(() => {{}});
      }} else if (vid.requestPictureInPicture) {{
        vid.requestPictureInPicture().catch(() => {{}});
      }}
    }}

    function toggleFullScreen() {{
      if (vid.requestFullscreen) {{
        vid.requestFullscreen();
      }}
    }}

    function captureSnapshot() {{
      const canvas = document.createElement('canvas');
      canvas.width = vid.videoWidth || 1080;
      canvas.height = vid.videoHeight || 1920;
      const ctx = canvas.getContext('2d');
      ctx.drawImage(vid, 0, 0, canvas.width, canvas.height);
      const link = document.createElement('a');
      link.download = `snapshot_frame_${{Math.round(vid.currentTime * FPS)}}.png`;
      link.href = canvas.toDataURL('image/png');
      link.click();
    }}

    window.addEventListener('keydown', (e) => {{
      if (e.target.tagName === 'INPUT' || e.target.tagName === 'SELECT') return;
      if (e.code === 'Space') {{ e.preventDefault(); togglePlay(); }}
      else if (e.code === 'ArrowLeft') {{ e.preventDefault(); vid.currentTime = Math.max(0, vid.currentTime - 3); }}
      else if (e.code === 'ArrowRight') {{ e.preventDefault(); vid.currentTime = Math.min(vid.duration, vid.currentTime + 3); }}
      else if (e.key === '[' || e.key === ',') {{ e.preventDefault(); stepFrame(-1); }}
      else if (e.key === ']' || e.key === '.') {{ e.preventDefault(); stepFrame(1); }}
      else if (e.key === 'f' || e.key === 'F') {{ e.preventDefault(); toggleFullScreen(); }}
      else if (e.key === 'm' || e.key === 'M') {{ e.preventDefault(); toggleMute(); }}
    }});
  </script>
</body>
</html>"""

    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as f:
        f.write(styled_html)
    return {"output_file": output_file.replace("\\", "/")}

def open_native(file_path: str):
    p = os.path.abspath(file_path)
    if os.name == 'nt':
        subprocess.Popen(f'cmd.exe /c start "" "{p}"', shell=True)
    elif sys.platform == 'darwin':
        subprocess.Popen(['open', p])
    else:
        subprocess.Popen(['xdg-open', p])

def main():
    parser = argparse.ArgumentParser(description="Universal Document & Video Previewer for Antigravity 2.0 & Windows")
    parser.add_argument("input_file", help="Path to PDF, DOCX, PPTX, XLSX, or Video (MP4/MOV/WEBM/MKV) file")
    parser.add_argument("-o", "--output", required=False, help="Output path (dir for PDF, file for HTML)")
    parser.add_argument("--dpi", type=int, default=150, help="DPI for PDF rendering")
    parser.add_argument("--open", action="store_true", help="Launch file directly with native OS application")
    args = parser.parse_args()

    p = Path(args.input_file)
    if not p.exists():
        print(f"File not found: {p}", file=sys.stderr)
        sys.exit(1)

    if args.open:
        open_native(str(p))
        print(f"LAUNCHED: Native application opened for {p}")

    suf = p.suffix.lower()
    out = args.output
    if not out:
        out = str(p.with_suffix('.html'))

    if suf == ".pdf":
        out_dir = out if not out.endswith('.html') else str(p.parent / f"{p.stem}_pdf_pages")
        r = preview_pdf(str(p), out_dir, args.dpi)
        print(f"SUCCESS: PDF {r['total_pages']} pages -> {r['output_dir']}")
    elif suf in [".docx", ".doc"]:
        out_file = out if out.endswith('.html') else os.path.join(out, f"{p.stem}.html")
        r = preview_docx(str(p), out_file)
        print(f"SUCCESS: DOCX -> {r['output_file']}")
    elif suf in [".pptx", ".ppt"]:
        out_file = out if out.endswith('.html') else os.path.join(out, f"{p.stem}.html")
        r = preview_pptx(str(p), out_file)
        print(f"SUCCESS: PPTX -> {r['output_file']}")
    elif suf in [".xlsx", ".xls", ".csv"]:
        out_file = out if out.endswith('.html') else os.path.join(out, f"{p.stem}.html")
        r = preview_xlsx(str(p), out_file)
        print(f"SUCCESS: XLSX -> {r['output_file']}")
    elif suf in [".mp4", ".mov", ".webm", ".mkv", ".m4v"]:
        out_file = out if out.endswith('.html') else os.path.join(out, f"{p.stem}.html")
        r = preview_video(str(p), out_file)
        print(f"SUCCESS: VIDEO -> {r['output_file']}")
    else:
        print(f"Unsupported format: {suf}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()


